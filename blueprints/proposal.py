"""
blueprints/proposal.py — Proposal routes extracted from server.py.

Flask Blueprint for all /api/proposals/* and /api/donor-templates endpoints.

Register in server.py with:
    from blueprints.proposal import proposal_bp
    app.register_blueprint(proposal_bp)
"""

import json
import logging
import os
import uuid

from flask import Blueprint, Response, jsonify, request

from auth import current_role, current_uid, require_auth, require_role
from blueprints.helpers import _chats_db, _log_event, _get_agent, _check_rate_limit

logger = logging.getLogger(__name__)

proposal_bp = Blueprint('proposal', __name__, url_prefix='/api')

# ── Per-user busy flag for advisor chat (prevents concurrent requests) ──────
import threading as _threading
_advisor_busy = {}
_advisor_busy_lock = _threading.Lock()
_ADVISOR_BUSY_TIMEOUT = 120  # seconds – auto-clear stuck flags


# ── Module-level time ref (same as server.py: import time as _time) ────────
import time as _time


# =============================================================================
# Proposal constants
# =============================================================================

PROPOSAL_SECTIONS = [
    "cover", "background", "needs_assessment", "toc", "logframe",
    "methodology", "budget", "mne_framework", "risk_matrix",
    "sustainability", "coordination", "final_review",
]

PROPOSAL_SECTION_LABELS = {
    "cover": "Cover Page",
    "background": "Context & Background",
    "needs_assessment": "Needs Assessment",
    "toc": "Theory of Change",
    "logframe": "Logical Framework",
    "methodology": "Methodology",
    "budget": "Budget Summary",
    "mne_framework": "Monitoring & Evaluation",
    "risk_matrix": "Risk Matrix",
    "sustainability": "Sustainability & Exit",
    "coordination": "Coordination",
    "final_review": "Final Review & Export",
}

SECTION_DB_FIELDS = {
    "cover": "cover_page",
    "background": "background",
    "needs_assessment": "needs_assessment",
    "toc": "toc",
    "logframe": "logframe",
    "methodology": "methodology",
    "budget": "budget",
    "mne_framework": "mne_framework",
    "risk_matrix": "risk_matrix",
    "sustainability": "sustainability",
    "coordination": "coordination",
    "final_review": "narrative",
}


# =============================================================================
# Helper functions
# =============================================================================

def _get_proposal_for_edit(prop_id: str, uid: str, role: str):
    """Fetch proposal row, check edit permissions. Returns (row, conn) or (None, conn)."""
    conn = _chats_db()
    try:
        if role == "admin":
            row = conn.execute("SELECT * FROM proposals WHERE id = ?", (prop_id,)).fetchone()
        else:
            row = conn.execute(
                "SELECT * FROM proposals WHERE id = ? AND uid = ?",
                (prop_id, uid)
            ).fetchone()
        return row, conn
    except Exception:
        conn.close()
        return None, None


def _update_step_status(conn, prop_id: str, step: str, status: str, uid: str, role: str):
    """Update the step_status JSON for a proposal."""
    row = conn.execute("SELECT step_status FROM proposals WHERE id = ?", (prop_id,)).fetchone()
    if not row:
        return
    try:
        step_status = json.loads(row["step_status"]) if row["step_status"] else {}
    except Exception:
        step_status = {}
    step_status[step] = status
    if role == "admin":
        conn.execute("UPDATE proposals SET step_status = ? WHERE id = ?", (json.dumps(step_status), prop_id))
    else:
        conn.execute("UPDATE proposals SET step_status = ? WHERE id = ? AND uid = ?", (json.dumps(step_status), prop_id, uid))
    conn.commit()


# =============================================================================
# Proposal CRUD routes
# =============================================================================

@proposal_bp.route("/proposals", methods=["GET"])
@require_auth
def api_get_proposals():
    uid = current_uid()
    role = current_role()
    conn = _chats_db()
    try:
        if role == "free":
            rows = conn.execute(
                """SELECT id, title, country, event, themes, donor, date_from, date_to,
                          current_step, step_status, created_at, completed_at
                   FROM proposals
                   WHERE uid = ? OR completed_at IS NOT NULL
                   ORDER BY created_at DESC""",
                (uid,)
            ).fetchall()
        else:
            rows = conn.execute(
                """SELECT id, title, country, event, themes, donor, date_from, date_to,
                          current_step, step_status, created_at, completed_at
                   FROM proposals
                   WHERE uid = ?
                   ORDER BY created_at DESC""",
                (uid,)
            ).fetchall()

        proposals = []
        for r in rows:
            try:
                themes_list = json.loads(r["themes"])
            except Exception:
                themes_list = [t.strip() for t in r["themes"].split(",") if t.strip()]

            try:
                step_status = json.loads(r["step_status"]) if r["step_status"] else {}
            except Exception:
                step_status = {}

            proposals.append({
                "id": r["id"],
                "title": r["title"],
                "country": r["country"],
                "event": r["event"],
                "themes": themes_list,
                "donor": r["donor"],
                "date_from": r["date_from"],
                "date_to": r["date_to"],
                "current_step": r["current_step"] or "cover",
                "step_status": step_status,
                "created_at": r["created_at"],
                "completed_at": r["completed_at"],
                "is_owner": r["id"] and uid and True or False,
            })
        return jsonify(proposals)
    except Exception as e:
        logger.error(f"api_get_proposals error: {e}")
        return jsonify({"error": "Internal server error"}), 500
    finally:
        conn.close()


@proposal_bp.route("/proposals/new", methods=["POST"])
@require_role("premium")
def api_create_proposal():
    uid = current_uid()
    role = current_role()
    data = request.json or {}

    title = data.get("title", "New Proposal").strip()
    country = data.get("country", "").strip()
    event = data.get("event", "").strip()
    themes = data.get("themes", [])
    donor = data.get("donor", "ECHO").strip()
    date_from = data.get("date_from", "").strip()
    date_to = data.get("date_to", "").strip()
    briefing = data.get("briefing", "").strip()

    # If briefing provided, merge it with any reference_text later
    briefing_text = ""
    if briefing:
        briefing_text = f"--- PROJECT BRIEFING (user-provided) ---\n{briefing}"

    if not country:
        return jsonify({"error": "Country context is required"}), 400

    conn = _chats_db()
    try:
        if role == "premium":
            monthly_count = conn.execute(
                "SELECT COUNT(*) FROM proposals WHERE uid = ? AND created_at > ?",
                (uid, _time.time() - 30 * 86400)
            ).fetchone()[0]
            if monthly_count >= 1:
                return jsonify({
                    "error": "Monthly proposal limit reached (1/month for premium). "
                             "Delete an existing proposal or upgrade to admin for unlimited.",
                    "limit": 1,
                    "used": monthly_count,
                    "premium_limit": True,
                }), 429

        prop_id = "prop_" + str(uuid.uuid4().hex[:12])

        default_toc = [
            {"level": "impact", "text": "Enhanced safety and reduced vulnerability of affected populations."},
            {"level": "outcome", "text": "Access to vital emergency services and basic needs is restored."},
            {"level": "output", "text": "Emergency relief kits and support materials distributed."},
            {"level": "activity", "text": "Procure and deliver aid packages to targeted zones."}
        ]

        default_logframe = {
            "goal": f"G1. Reduced vulnerability to disaster shocks in {country}.",
            "outcomes": "OC1. Targeted households report basic needs met.\nIndicator: % of target pop with satisfied needs.",
            "outputs": "O1. Relief materials delivered to local centers.\nIndicator: Number of kits distributed.",
            "activities": "A1. Deploy logistics team.\nA2. Complete safe distributions."
        }

        default_narrative = f"## Project Summary\nEmergency humanitarian response targeting communities in {country} affected by recent crises.\n\n## Methodology\nInterventions will focus on key sectors: {', '.join(themes)}."

        default_step_status = {step: "locked" for step in [
            "cover", "background", "needs_assessment", "toc", "logframe",
            "methodology", "budget", "mne_framework", "risk_matrix",
            "sustainability", "coordination", "final_review"
        ]}
        default_step_status["cover"] = "empty"

        conn.execute(
            """INSERT INTO proposals
               (id, uid, title, country, event, themes, donor, date_from, date_to,
                toc, logframe, narrative, created_at,
                cover_page, background, needs_assessment, methodology, budget,
                mne_framework, risk_matrix, sustainability, coordination,
                current_step, step_status, completed_at, reference_text, reference_filename)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, '{}', '', '', '', '{}', '{}', '[]', '', '', 'cover', ?, NULL, ?, '')""",
            (
                prop_id, uid, title, country, event,
                json.dumps(themes), donor, date_from, date_to,
                json.dumps(default_toc), json.dumps(default_logframe), default_narrative,
                _time.time(),
                json.dumps(default_step_status),
                briefing_text,
            )
        )
        conn.commit()
        _log_event(uid, "proposal_created", {"prop_id": prop_id, "role": role})

        return jsonify({
            "id": prop_id,
            "title": title,
            "country": country,
            "event": event,
            "themes": themes,
            "donor": donor,
            "date_from": date_from,
            "date_to": date_to,
            "toc": default_toc,
            "logframe": default_logframe,
            "narrative": default_narrative,
            "cover_page": {},
            "background": "",
            "needs_assessment": "",
            "methodology": "",
            "budget": {},
            "mne_framework": {},
            "risk_matrix": [],
            "sustainability": "",
            "coordination": "",
            "current_step": "cover",
            "step_status": default_step_status,
            "completed_at": None,
            "has_reference": bool(briefing_text),
            "reference_filename": "",
        }), 201
    except Exception as e:
        logger.error(f"api_create_proposal error: {e}")
        return jsonify({"error": "Internal server error"}), 500
    finally:
        conn.close()


@proposal_bp.route("/proposals/<prop_id>", methods=["GET"])
@require_auth
def api_get_proposal_detail(prop_id):
    uid = current_uid()
    role = current_role()
    conn = _chats_db()
    try:
        row = conn.execute(
            "SELECT * FROM proposals WHERE id = ?",
            (prop_id,)
        ).fetchone()

        if not row:
            return jsonify({"error": "Proposal not found"}), 404

        is_owner = row["uid"] == uid
        is_admin = role == "admin"
        is_completed = row["completed_at"] is not None

        if not is_owner and not is_admin:
            if role == "free" and not is_completed:
                return jsonify({"error": "This proposal is not yet published", "premium_required": False}), 403
            if role == "premium" and not is_completed:
                return jsonify({"error": "You can only view completed proposals from other users"}), 403

        can_edit = (is_owner and role in ("premium", "admin")) or is_admin

        try:
            themes_list = json.loads(row["themes"])
        except Exception:
            themes_list = [t.strip() for t in row["themes"].split(",") if t.strip()]

        try:
            step_status = json.loads(row["step_status"]) if row["step_status"] else {}
        except Exception:
            step_status = {}

        return jsonify({
            "id": row["id"],
            "title": row["title"],
            "country": row["country"],
            "event": row["event"],
            "themes": themes_list,
            "donor": row["donor"],
            "date_from": row["date_from"],
            "date_to": row["date_to"],
            "toc": json.loads(row["toc"]),
            "logframe": json.loads(row["logframe"]),
            "narrative": row["narrative"],
            "cover_page": json.loads(row["cover_page"]) if row["cover_page"] else {},
            "background": row["background"] or "",
            "needs_assessment": row["needs_assessment"] or "",
            "methodology": row["methodology"] or "",
            "budget": json.loads(row["budget"]) if row["budget"] else {},
            "mne_framework": json.loads(row["mne_framework"]) if row["mne_framework"] else {},
            "risk_matrix": json.loads(row["risk_matrix"]) if row["risk_matrix"] else [],
            "sustainability": row["sustainability"] or "",
            "coordination": row["coordination"] or "",
            "current_step": row["current_step"] or "cover",
            "step_status": step_status,
            "created_at": row["created_at"],
            "completed_at": row["completed_at"],
            "can_edit": can_edit,
            "is_owner": is_owner,
            "reference_filename": row["reference_filename"] if "reference_filename" in row.keys() else "",
            "has_reference": bool(row["reference_text"]) if "reference_text" in row.keys() else False,
        })
    except Exception as e:
        logger.error(f"api_get_proposal_detail error: {prop_id}, {e}")
        return jsonify({"error": "Internal server error"}), 500
    finally:
        conn.close()


@proposal_bp.route("/proposals/<prop_id>", methods=["PUT"])
@require_role("premium")
def api_update_proposal(prop_id):
    uid = current_uid()
    role = current_role()
    data = request.json or {}

    conn = _chats_db()
    try:
        if role == "admin":
            row = conn.execute(
                "SELECT id FROM proposals WHERE id = ?", (prop_id,)
            ).fetchone()
        else:
            row = conn.execute(
                "SELECT id FROM proposals WHERE id = ? AND uid = ?",
                (prop_id, uid)
            ).fetchone()

        if not row:
            return jsonify({"error": "Proposal not found"}), 404

        allowed_fields = [
            "title", "country", "event", "donor", "date_from", "date_to",
            "narrative", "background", "needs_assessment", "methodology",
            "sustainability", "coordination", "reference_text",
        ]
        json_fields = ["themes", "toc", "logframe", "cover_page", "budget", "mne_framework", "risk_matrix"]

        fields_to_update = {}
        for k in allowed_fields:
            if k in data:
                fields_to_update[k] = data[k]

        for k in json_fields:
            if k in data:
                fields_to_update[k] = json.dumps(data[k])

        if "current_step" in data:
            fields_to_update["current_step"] = data["current_step"]
        if "step_status" in data:
            fields_to_update["step_status"] = json.dumps(data["step_status"])

        # Validate field names against whitelist (prevent SQL injection)
        set_parts = []
        params = []
        for k in allowed_fields:
            if k in data:
                set_parts.append(f"{k} = ?")
                params.append(data[k])
        for k in json_fields:
            if k in data:
                set_parts.append(f"{k} = ?")
                params.append(json.dumps(data[k]))

        if not set_parts:
            return jsonify({"message": "No changes made"})

        set_clause = ", ".join(set_parts)

        if role == "admin":
            params.append(prop_id)
            conn.execute(
                f"UPDATE proposals SET {set_clause} WHERE id = ?",
                params
            )
        else:
            params.extend([prop_id, uid])
            conn.execute(
                f"UPDATE proposals SET {set_clause} WHERE id = ? AND uid = ?",
                params
            )
        conn.commit()
        return jsonify({"message": "Proposal updated successfully"})
    except Exception as e:
        logger.error(f"api_update_proposal error: {prop_id}, {e}")
        return jsonify({"error": "Internal server error"}), 500
    finally:
        conn.close()


@proposal_bp.route("/proposals/<prop_id>", methods=["DELETE"])
@require_role("premium")
def api_delete_proposal(prop_id):
    uid = current_uid()
    role = current_role()
    conn = _chats_db()
    try:
        if role == "admin":
            row = conn.execute("SELECT id FROM proposals WHERE id = ?", (prop_id,)).fetchone()
        else:
            row = conn.execute(
                "SELECT id FROM proposals WHERE id = ? AND uid = ?",
                (prop_id, uid)
            ).fetchone()

        if not row:
            return jsonify({"error": "Proposal not found"}), 404

        if role == "admin":
            conn.execute("DELETE FROM proposals WHERE id = ?", (prop_id,))
        else:
            conn.execute("DELETE FROM proposals WHERE id = ? AND uid = ?", (prop_id, uid))
        conn.commit()
        _log_event(uid, "proposal_deleted", {"prop_id": prop_id})
        return jsonify({"message": "Proposal deleted successfully"})
    except Exception as e:
        logger.error(f"api_delete_proposal error: {prop_id}, {e}")
        return jsonify({"error": "Internal server error"}), 500
    finally:
        conn.close()


# =============================================================================
# Pinned sources
# =============================================================================

@proposal_bp.route("/proposals/<prop_id>/pin-source", methods=["POST"])
@require_auth
def api_pin_source(prop_id):
    uid = current_uid()
    data = request.json or {}
    conn = _chats_db()
    try:
        row = conn.execute("SELECT id, pinned_sources FROM proposals WHERE id = ? AND uid = ?", (prop_id, uid)).fetchone()
        if not row: return jsonify({"error": "Proposal not found"}), 404

        pinned = []
        if row["pinned_sources"]:
            try: pinned = json.loads(row["pinned_sources"])
            except (json.JSONDecodeError, TypeError): pass  # json parse

        pinned.append({
            "id": data.get("id", str(uuid.uuid4())),
            "title": data.get("title", ""),
            "url": data.get("url", ""),
            "snippet": data.get("snippet", "")
        })

        conn.execute("UPDATE proposals SET pinned_sources = ? WHERE id = ?", (json.dumps(pinned), prop_id))
        conn.commit()
        return jsonify({"status": "success", "pinned_sources": pinned})
    finally:
        conn.close()


@proposal_bp.route("/proposals/<prop_id>/pin-source/<int:index>", methods=["DELETE"])
@require_auth
def api_delete_pinned_source(prop_id, index):
    uid = current_uid()
    conn = _chats_db()
    try:
        row = conn.execute("SELECT id, pinned_sources FROM proposals WHERE id = ? AND uid = ?", (prop_id, uid)).fetchone()
        if not row: return jsonify({"error": "Proposal not found"}), 404

        pinned = []
        if row["pinned_sources"]:
            try: pinned = json.loads(row["pinned_sources"])
            except (json.JSONDecodeError, TypeError): pass  # json parse

        if 0 <= index < len(pinned):
            pinned.pop(index)
            conn.execute("UPDATE proposals SET pinned_sources = ? WHERE id = ?", (json.dumps(pinned), prop_id))
            conn.commit()
            return jsonify({"status": "success", "pinned_sources": pinned})
        else:
            return jsonify({"error": "Index out of bounds"}), 400
    finally:
        conn.close()


# =============================================================================
# ToC & Logframe updates
# =============================================================================

@proposal_bp.route("/proposals/<prop_id>/toc", methods=["PUT"])
@require_auth
def api_update_toc(prop_id):
    uid = current_uid()
    data = request.json or {}
    conn = _chats_db()
    try:
        row = conn.execute("SELECT id FROM proposals WHERE id = ? AND uid = ?", (prop_id, uid)).fetchone()
        if not row: return jsonify({"error": "Proposal not found"}), 404
        nodes = data.get("nodes", [])
        conn.execute("UPDATE proposals SET toc_nodes = ? WHERE id = ?", (json.dumps(nodes), prop_id))
        conn.commit()
        return jsonify({"message": "Theory of Change updated", "toc_nodes": nodes})
    finally:
        conn.close()


@proposal_bp.route("/proposals/<prop_id>/logframe", methods=["PUT"])
@require_auth
def api_update_logframe(prop_id):
    uid = current_uid()
    data = request.json or {}
    conn = _chats_db()
    try:
        row = conn.execute("SELECT id, logframe_data FROM proposals WHERE id = ? AND uid = ?", (prop_id, uid)).fetchone()
        if not row: return jsonify({"error": "Proposal not found"}), 404

        lf = {}
        if row["logframe_data"]:
            try: lf = json.loads(row["logframe_data"])
            except (json.JSONDecodeError, TypeError): pass  # json parse

        section = data.get("section")
        index = data.get("index")
        field = data.get("field")
        value = data.get("value")

        if section and field and index is not None:
            if section not in lf: lf[section] = []
            while len(lf[section]) <= index:
                lf[section].append({})
            lf[section][index][field] = value

            conn.execute("UPDATE proposals SET logframe_data = ? WHERE id = ?", (json.dumps(lf), prop_id))
            conn.commit()
            return jsonify({"message": "Logframe cell updated", "logframe_data": lf})
        else:
            if "logframe_data" in data:
                conn.execute("UPDATE proposals SET logframe_data = ? WHERE id = ?", (json.dumps(data["logframe_data"]), prop_id))
                conn.commit()
                return jsonify({"message": "Logframe updated", "logframe_data": data["logframe_data"]})
            return jsonify({"error": "Invalid payload"}), 400
    finally:
        conn.close()


@proposal_bp.route("/proposals/<prop_id>/budget-calc", methods=["PUT"])
@require_auth
def api_budget_calc(prop_id):
    uid = current_uid()
    data = request.json or {}
    conn = _chats_db()
    try:
        row = conn.execute("SELECT id, budget_details FROM proposals WHERE id = ? AND uid = ?", (prop_id, uid)).fetchone()
        if not row: return jsonify({"error": "Proposal not found"}), 404

        lines = data.get("lines", [])
        # Validate lines structure
        if not isinstance(lines, list):
            return jsonify({"error": "lines must be an array"}), 400
        if len(lines) > 200:
            return jsonify({"error": "Too many budget lines (max 200)"}), 400
        for line in lines:
            if not isinstance(line, dict):
                return jsonify({"error": "Each line must be an object"}), 400
            if "category" not in line or "amount" not in line:
                return jsonify({"error": "Each line must have 'category' and 'amount'"}), 400
            try:
                amt = float(line.get("amount", 0))
                if amt < 0:
                    return jsonify({"error": "Amount cannot be negative"}), 400
            except (ValueError, TypeError):
                return jsonify({"error": "Amount must be a number"}), 400

        total = sum(float(line.get("amount", 0)) for line in lines)
        percentages = {}
        if total > 0:
            for line in lines:
                cat = line.get("category", "Unknown")
                amt = float(line.get("amount", 0))
                percentages[cat] = f"{(amt / total) * 100:.1f}%"

        donor_limit_ok = True
        budget_det = {
            "lines": lines,
            "total": total,
            "percentages": percentages,
            "donor_limit_ok": donor_limit_ok
        }

        conn.execute("UPDATE proposals SET budget_details = ? WHERE id = ?", (json.dumps(budget_det), prop_id))
        conn.commit()
        return jsonify(budget_det)
    finally:
        conn.close()


# =============================================================================
# Apply suggestion
# =============================================================================

@proposal_bp.route("/proposals/<prop_id>/apply-suggestion", methods=["POST"])
@require_auth
def api_apply_suggestion(prop_id):
    uid = current_uid()
    data = request.json or {}
    action = data.get("action")
    payload = data.get("payload", {})

    conn = _chats_db()
    try:
        row = conn.execute("SELECT * FROM proposals WHERE id = ? AND uid = ?", (prop_id, uid)).fetchone()
        if not row: return jsonify({"error": "Proposal not found"}), 404

        if action == "update_logframe":
            lf = {}
            if row["logframe_data"]:
                try: lf = json.loads(row["logframe_data"])
                except (json.JSONDecodeError, TypeError): pass  # json parse
            sec, idx, fld, val = payload.get("section"), payload.get("index"), payload.get("field"), payload.get("value")
            if sec and fld and idx is not None:
                if sec not in lf: lf[sec] = []
                while len(lf[sec]) <= idx: lf[sec].append({})
                lf[sec][idx][fld] = val
                conn.execute("UPDATE proposals SET logframe_data = ? WHERE id = ?", (json.dumps(lf), prop_id))
                conn.commit()
                return jsonify({"status": "applied", "new_state": lf})

        return jsonify({"error": "Unsupported action"}), 400
    finally:
        conn.close()


# =============================================================================
# Admin proposal delete
# =============================================================================

@proposal_bp.route("/admin/proposals/<prop_id>", methods=["DELETE"])
@require_role("admin")
def api_admin_delete_proposal(prop_id):
    """Delete any proposal regardless of owner (admin only)."""
    conn = _chats_db()
    try:
        row = conn.execute("SELECT id FROM proposals WHERE id = ?", (prop_id,)).fetchone()
        if not row:
            return jsonify({"error": "Proposal not found"}), 404
        conn.execute("DELETE FROM proposals WHERE id = ?", (prop_id,))
        conn.commit()
        _log_event(current_uid(), "proposal_deleted", {"prop_id": prop_id, "admin": True})
        return jsonify({"message": "Proposal deleted"})
    except Exception as e:
        logger.error(f"api_admin_delete_proposal error: {prop_id}, {e}")
        return jsonify({"error": "Internal server error"}), 500
    finally:
        conn.close()


# =============================================================================
# Generate ToC / Logframe / Narrative (LLM)
# =============================================================================

@proposal_bp.route("/proposals/<prop_id>/generate-toc", methods=["POST"])
@require_auth
def api_proposal_generate_toc(prop_id):
    uid = current_uid()
    conn = _chats_db()
    try:
        row = conn.execute(
            "SELECT country, event, themes FROM proposals WHERE id = ? AND uid = ?",
            (prop_id, uid)
        ).fetchone()

        if not row:
            return jsonify({"error": "Proposal not found"}), 404

        country = row["country"]
        event = row["event"]
        try:
            themes = json.loads(row["themes"])
        except Exception:
            themes = [row["themes"]]

        context_chunks = []
        try:
            from reliefweb_api.vector_store import VectorStore
            store = VectorStore()
            query = f"{country} {event} {' '.join(themes)}"
            results = store.search(query=query, limit=5, country=country)
            for res in results.get("results", []):
                context_chunks.append(res.get("text", ""))
        except BaseException as vec_err:
            logger.warning(f"Vector search failed in generate-toc: {vec_err}")

        context_text = "\n\n".join(context_chunks)[:4000]

        system_prompt = (
            "You are an expert humanitarian crisis proposal designer.\n"
            "Your task is to draft a Theory of Change (ToC) for a relief project.\n"
            "Analyze the provided crisis context and return a structured JSON array representing the ToC levels.\n"
            "The JSON array MUST contain exactly 4 objects corresponding to the standard ToC levels in this order:\n"
            "1. Goal / Impact (Ultimate long-term change)\n"
            "2. Outcome (Specific change in behavior/status of target pop)\n"
            "3. Output (Direct product of project activities)\n"
            "4. Activity (Key action to produce outputs)\n\n"
            "Each object must have two fields: 'level' ('impact', 'outcome', 'output', 'activity') and 'text'.\n"
            "Ensure the logic flows sequentially (Activity -> Output -> Outcome -> Impact).\n"
            "Return ONLY the JSON array, no explanation or markdown blocks."
        )

        user_prompt = (
            f"Country: {country}\n"
            f"Crisis / Event: {event}\n"
            f"Target Themes: {', '.join(themes)}\n\n"
            f"Crisis Context Data:\n{context_text if context_text else 'No recent report details available.'}"
        )

        from sitrep.llm_client import chat as llm_chat
        response = llm_chat([
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt}
        ])

        from sitrep.utils import clean_json_response
        cleaned = clean_json_response(response)

        toc_nodes = json.loads(cleaned)

        conn.execute(
            "UPDATE proposals SET toc = ? WHERE id = ? AND uid = ?",
            (json.dumps(toc_nodes), prop_id, uid)
        )
        conn.commit()

        return jsonify(toc_nodes)
    except Exception as e:
        logger.error(f"api_proposal_generate_toc error: {prop_id}, {e}")
        return jsonify({"error": "Generation failed. Please try again."}), 500
    finally:
        conn.close()


@proposal_bp.route("/proposals/<prop_id>/generate-logframe", methods=["POST"])
@require_auth
def api_proposal_generate_logframe(prop_id):
    uid = current_uid()
    conn = _chats_db()
    try:
        row = conn.execute(
            "SELECT country, event, themes, toc FROM proposals WHERE id = ? AND uid = ?",
            (prop_id, uid)
        ).fetchone()

        if not row:
            return jsonify({"error": "Proposal not found"}), 404

        country = row["country"]
        event = row["event"]
        toc = json.loads(row["toc"])

        system_prompt = (
            "You are an expert humanitarian program officer.\n"
            "Based on the Theory of Change (ToC) provided, draft a structured Logical Framework (Logframe) matrix.\n"
            "Return a JSON object containing the primary logic sections:\n"
            "- 'goal': Statement of impact + key indicator\n"
            "- 'outcomes': Statement of outcome + key indicator\n"
            "- 'outputs': Direct outputs + key indicators\n"
            "- 'activities': Direct activities.\n"
            "Make all indicators SMART (Specific, Measurable, Achievable, Relevant, Time-bound).\n"
            "Return ONLY the JSON object, no explanation or markdown blocks."
        )

        user_prompt = (
            f"Country: {country}\n"
            f"Crisis: {event}\n"
            f"Theory of Change Hierarchy:\n" +
            "\n".join([f"- {node['level'].upper()}: {node['text']}" for node in toc])
        )

        from sitrep.llm_client import chat as llm_chat
        response = llm_chat([
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt}
        ])

        from sitrep.utils import clean_json_response
        cleaned = clean_json_response(response)

        logframe_data = json.loads(cleaned)

        conn.execute(
            "UPDATE proposals SET logframe = ? WHERE id = ? AND uid = ?",
            (json.dumps(logframe_data), prop_id, uid)
        )
        conn.commit()

        return jsonify(logframe_data)
    except Exception as e:
        logger.error(f"api_proposal_generate_logframe error: {prop_id}, {e}")
        return jsonify({"error": "Generation failed. Please try again."}), 500
    finally:
        conn.close()


@proposal_bp.route("/proposals/<prop_id>/generate-narrative", methods=["POST"])
@require_role("premium")
def api_proposal_generate_narrative(prop_id):
    uid = current_uid()
    conn = _chats_db()
    try:
        row = conn.execute(
            "SELECT country, event, donor, toc, logframe FROM proposals WHERE id = ? AND uid = ?",
            (prop_id, uid)
        ).fetchone()

        if not row:
            return jsonify({"error": "Proposal not found"}), 404

        country = row["country"]
        event = row["event"]
        donor = row["donor"]
        toc = row["toc"]
        logframe = row["logframe"]

        system_prompt = (
            f"You are a professional grant proposal writer specialized in {donor} application guidelines.\n"
            f"Draft the full project description narrative matching {donor} standard templates.\n"
            "Use clear Markdown formatting with headers.\n"
            "Structure it into: 1. Needs Assessment, 2. Project Description, 3. Logical Framework Alignment, 4. Sustainability & Risks.\n"
            "Incorporate details from the Logical Framework and Theory of Change provided.\n"
            "Maintain a formal, data-driven, and highly persuasive tone. Do not write placeholders."
        )

        user_prompt = (
            f"Crisis Context: {country} / {event}\n"
            f"Theory of Change: {toc}\n"
            f"Logical Framework Matrix: {logframe}"
        )

        from sitrep.llm_client import chat as llm_chat
        response = llm_chat([
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt}
        ])

        conn.execute(
            "UPDATE proposals SET narrative = ? WHERE id = ? AND uid = ?",
            (response, prop_id, uid)
        )
        conn.commit()

        return jsonify({"narrative": response})
    except Exception as e:
        logger.error(f"api_proposal_generate_narrative error: {prop_id}, {e}")
        return jsonify({"error": "Narrative generation failed. Please try again."}), 500
    finally:
        conn.close()


# =============================================================================
# Proposal chunks (RAG)
# =============================================================================

@proposal_bp.route("/proposals/<prop_id>/chunks", methods=["GET"])
@require_auth
def api_proposal_chunks(prop_id):
    uid = current_uid()
    conn = _chats_db()
    try:
        row = conn.execute(
            "SELECT country, themes, date_from, date_to FROM proposals WHERE id = ? AND uid = ?",
            (prop_id, uid)
        ).fetchone()

        if not row:
            return jsonify({"error": "Proposal not found"}), 404

        country = row["country"]
        try:
            themes = json.loads(row["themes"])
        except Exception:
            themes = [t.strip() for t in row["themes"].split(",") if t.strip()]

        date_from = row["date_from"]
        date_to = row["date_to"]

        chunks = []
        try:
            from sitrep.chroma_adapter import ChromaAdapter
            db = ChromaAdapter()
            chunks = db.get_chunks_by_country_and_themes(
                country, themes or None, date_from=date_from or None, date_to=date_to or None,
            )
        except BaseException as adapter_err:
            logger.warning(f"ChromaAdapter failed in api_proposal_chunks: {adapter_err}")

        results = []
        for c in chunks[:15]:
            results.append({
                "text": c.get("text", ""),
                "title": c.get("title", "Situation Report"),
                "date": c.get("date", ""),
                "themes": c.get("themes", "")
            })

        return jsonify(results)
    except Exception as e:
        logger.error(f"api_proposal_chunks error: {prop_id}, {e}")
        return jsonify([])
    finally:
        conn.close()


# =============================================================================
# Proposal Advisor chat
# =============================================================================

@proposal_bp.route("/proposals/<prop_id>/advisor/chat", methods=["POST"])
@require_auth
def api_proposal_advisor_chat(prop_id):
    uid = current_uid()
    role = current_role()
    data = request.json or {}
    message = data.get("message", "").strip()

    if not message:
        return jsonify({"error": "Message is required"}), 400

    # ── Per-user busy flag (prevents concurrent advisor requests) ──────────
    with _advisor_busy_lock:
        # Auto-clear stuck busy flags (older than timeout)
        if uid in _advisor_busy and (_time.time() - _advisor_busy.get(uid + "_since", 0)) > _ADVISOR_BUSY_TIMEOUT:
            _advisor_busy[uid] = False
        if _advisor_busy.get(uid, False):
            return jsonify({"error": "Advisor is busy processing your previous request. Please wait."}), 429

        # Daily rate limit check
        rate = _check_rate_limit(uid, role)
        if not rate["allowed"]:
            _log_event(uid, "rate_limit_hit", {"reason": "daily_limit", "endpoint": "advisor_chat", "limit": rate["limit"]})
            return jsonify({"error": "Daily limit reached", "limit": rate["limit"], "used": rate["used"]}), 429

        _advisor_busy[uid] = True
        _advisor_busy[uid + "_since"] = _time.time()

    try:
        conn = _chats_db()
        # Verify proposal ownership
        row = conn.execute(
            "SELECT country, event, donor, toc, logframe FROM proposals WHERE id = ? AND uid = ?",
            (prop_id, uid)
        ).fetchone()

        if not row:
            return jsonify({"error": "Proposal not found"}), 404

        chat_id = f"proposal_advisor_{prop_id}"

        # Ensure advisor chat session exists in chats table
        chat_row = conn.execute("SELECT id FROM chats WHERE id = ?", (chat_id,)).fetchone()
        if not chat_row:
            conn.execute(
                "INSERT INTO chats (id, uid, title, created) VALUES (?, ?, ?, ?)",
                (chat_id, uid, f"Advisor: {row['country']} Proposal", _time.time())
            )
            conn.commit()

        # Get historical advisor messages
        db_rows = conn.execute(
            "SELECT role, content FROM chat_messages WHERE chat_id = ? ORDER BY ts ASC",
            (chat_id,)
        ).fetchall()

        from langchain_core.messages import AIMessage, HumanMessage, ToolMessage
        messages = []
        for r in db_rows:
            if r["role"] == "user":
                messages.append(HumanMessage(content=r["content"]))
            elif r["role"] == "assistant":
                messages.append(AIMessage(content=r["content"]))

        # Append the new user message
        messages.append(HumanMessage(content=message))

        # Save user message to database
        conn.execute(
            "INSERT INTO chat_messages (chat_id, role, content, ts) VALUES (?, 'user', ?, ?)",
            (chat_id, message, _time.time())
        )
        conn.commit()

        # Fetch context chunks
        chunks_text = ""
        try:
            from sitrep.chroma_adapter import ChromaAdapter
            db = ChromaAdapter()
            try:
                themes_list = json.loads(row["themes"])
            except Exception:
                themes_list = [t.strip() for t in row["themes"].split(",") if t.strip()]

            # Use row data from proposal (country, themes) to get chunks
            chunks = db.get_chunks_by_country_and_themes(row["country"], themes_list or None)
            if chunks:
                chunks_text = "\n\n".join([f"- {c.get('title', 'Report')}: {c.get('text', '')}" for c in chunks[:10]])
        except Exception as e:
            logger.warning(f"Failed to fetch chunks for advisor: {e}")

        # Inject a SystemMessage with Proposal Context
        from langchain_core.messages import SystemMessage
        advisor_context = f"""
You are the Proposal Design Advisor. The user is actively working on a proposal.
Here is the current state of the proposal:
Country: {row['country']}
Event: {row['event']}
Donor: {row['donor']}

Theory of Change:
{row['toc']}

Logframe:
{row['logframe']}

Here is recent relevant background data (RAG context):
{chunks_text}

Provide specific, constructive feedback and suggestions. Use your tools (edit_proposal_toc, edit_proposal_logframe, edit_proposal_narrative) to apply changes directly when asked.
"""
        messages.insert(0, SystemMessage(content=advisor_context))

        # Invoke agent
        agent = _get_agent()
        config = {
            "recursion_limit": 25,
            "configurable": {
                "uid": uid,
                "proposal_id": prop_id
            }
        }

        result = agent.invoke({"messages": messages}, config=config)

        # Save agent response to database
        final_response = ""
        for msg in reversed(result.get("messages", [])):
            if isinstance(msg, AIMessage) and not getattr(msg, "tool_calls", None):
                final_response = msg.content
                break

        if not final_response:
            final_response = "I have reviewed your proposal details."

        conn.execute(
            "INSERT INTO chat_messages (chat_id, role, content, ts) VALUES (?, 'assistant', ?, ?)",
            (chat_id, final_response, _time.time())
        )
        conn.commit()

        # Check if proposal tools were executed to trigger auto-refresh in frontend
        proposal_edited = False
        for msg in result.get("messages", []):
            if isinstance(msg, ToolMessage) and msg.name in ("edit_proposal_toc", "edit_proposal_logframe", "edit_proposal_narrative"):
                proposal_edited = True
                break

        command_data = None
        if proposal_edited:
            command_data = {"action": "refresh"}

        # Legacy support for text command tags
        import re
        cmd_match = re.search(r"<cmd>(.*?)</cmd>", final_response, re.DOTALL)
        if cmd_match:
            try:
                command_data = json.loads(cmd_match.group(1).strip())
                final_response = final_response.replace(cmd_match.group(0), "").strip()

                # Apply legacy command to database
                if command_data.get("action") == "update_logframe":
                    field = command_data.get("field")
                    text_val = command_data.get("text")
                    parsed_lf = json.loads(row["logframe"])
                    if field in parsed_lf:
                        parsed_lf[field] = text_val
                        conn.execute(
                            "UPDATE proposals SET logframe = ? WHERE id = ? AND uid = ?",
                            (json.dumps(parsed_lf), prop_id, uid)
                        )
                        conn.commit()
                elif command_data.get("action") == "update_toc":
                    index = command_data.get("index")
                    text_val = command_data.get("text")
                    parsed_toc = json.loads(row["toc"])
                    if 0 <= index < len(parsed_toc):
                        parsed_toc[index]["text"] = text_val
                        conn.execute(
                            "UPDATE proposals SET toc = ? WHERE id = ? AND uid = ?",
                            (json.dumps(parsed_toc), prop_id, uid)
                        )
                        conn.commit()
            except Exception as parse_err:
                logger.warning(f"Failed to parse or apply advisor command JSON: {parse_err}")

        return jsonify({
            "response": final_response,
            "command": command_data
        })
    except Exception as e:
        logger.error("api_proposal_advisor_chat error: %s, %s", prop_id, e)
        return jsonify({"error": "Advisor failed. Please try again."}), 500
    finally:
        with _advisor_busy_lock:
            _advisor_busy.pop(uid, None)
            _advisor_busy.pop(uid + "_since", None)
        try:
            conn.close()
        except Exception:
            pass


@proposal_bp.route("/proposals/<prop_id>/advisor/background-review", methods=["POST"])
@require_auth
def api_proposal_advisor_background_review(prop_id):
    uid = current_uid()
    conn = _chats_db()
    try:
        row = conn.execute(
            "SELECT country, event, donor, toc, logframe, narrative, themes FROM proposals WHERE id = ? AND uid = ?",
            (prop_id, uid)
        ).fetchone()

        if not row:
            return jsonify({"error": "Proposal not found"}), 404

        # Fetch context chunks
        chunks_text = ""
        try:
            from sitrep.chroma_adapter import ChromaAdapter
            db = ChromaAdapter()
            try:
                themes_list = json.loads(row["themes"])
            except Exception:
                themes_list = [t.strip() for t in row["themes"].split(",") if t.strip()]

            chunks = db.get_chunks_by_country_and_themes(row["country"], themes_list or None)
            if chunks:
                chunks_text = "\n\n".join([f"- {c.get('title', 'Report')}: {c.get('text', '')}" for c in chunks[:10]])
        except Exception as e:
            logger.warning(f"Failed to fetch chunks for background review: {e}")

        from langchain_core.messages import AIMessage, HumanMessage, SystemMessage
        advisor_context = f"""
You are the Proposal Background AI Advisor. The user just saved a field update to their proposal.
Your task is to review the current proposal state in the background and suggest improvements.

Current Proposal:
Country: {row['country']}
Event: {row['event']}
Donor: {row['donor']}

Theory of Change:
{row['toc']}

Logframe:
{row['logframe']}

Narrative text:
{row['narrative']}

Here is recent relevant background data (RAG context):
{chunks_text}

CRITICAL: Do NOT use tools that directly edit the database (e.g., edit_proposal_toc, edit_proposal_logframe, edit_proposal_narrative).
Instead, MUST ONLY use the `propose_edits` tool if you want to suggest concrete changes to the ToC, Logframe, or Narrative. 
If everything looks perfect and no edits are needed, just reply with an encouraging message and do not call `propose_edits`.
"""
        messages = [
            SystemMessage(content=advisor_context),
            HumanMessage(content="Please review my recent updates and propose edits if necessary.")
        ]

        agent = _get_agent()
        config = {
            "recursion_limit": 25,
            "configurable": {
                "uid": uid,
                "proposal_id": prop_id
            }
        }

        result = agent.invoke({"messages": messages}, config=config)

        final_message = "Background review complete."
        drafts = {}

        for msg in result.get("messages", []):
            if isinstance(msg, AIMessage) and getattr(msg, "content", "") and not getattr(msg, "tool_calls", None):
                final_message = msg.content
            if isinstance(msg, AIMessage) and getattr(msg, "tool_calls", None):
                for call in msg.tool_calls:
                    if call["name"] == "propose_edits":
                        args = call["args"]
                        if args.get("toc"):
                            drafts["toc"] = args["toc"]
                        if args.get("logframe"):
                            drafts["logframe"] = args["logframe"]
                        if args.get("narrative"):
                            drafts["narrative"] = args["narrative"]

        # Log to chat history to show background action in Advisor
        chat_id = f"proposal_advisor_{prop_id}"
        chat_row = conn.execute("SELECT id FROM chats WHERE id = ?", (chat_id,)).fetchone()
        if chat_row:
            msg_to_save = final_message
            if drafts:
                msg_to_save = final_message + "\n\n*(I have prepared proposed drafts for your review. Click the Review button to see them.)*"

            conn.execute(
                "INSERT INTO chat_messages (chat_id, role, content, ts) VALUES (?, 'assistant', ?, ?)",
                (chat_id, msg_to_save, _time.time())
            )
            conn.commit()

        return jsonify({
            "message": final_message,
            "drafts": drafts if drafts else None
        })
    except Exception as e:
        logger.error(f"api_proposal_advisor_background_review error: {prop_id}, {e}")
        return jsonify({"error": f"Background review failed: {str(e)}"}), 500
    finally:
        conn.close()


@proposal_bp.route("/proposals/<prop_id>/advisor/history", methods=["GET"])
@require_auth
def api_proposal_advisor_history(prop_id):
    uid = current_uid()
    chat_id = f"proposal_advisor_{prop_id}"
    conn = _chats_db()
    try:
        # Check proposal ownership
        row = conn.execute(
            "SELECT id FROM proposals WHERE id = ? AND uid = ?",
            (prop_id, uid)
        ).fetchone()
        if not row:
            return jsonify({"error": "Proposal not found"}), 404

        db_rows = conn.execute(
            "SELECT role, content FROM chat_messages WHERE chat_id = ? ORDER BY ts ASC",
            (chat_id,)
        ).fetchall()

        history = [{"role": r["role"], "content": r["content"]} for r in db_rows]
        return jsonify(history)
    except Exception as e:
        logger.error(f"api_proposal_advisor_history error: {prop_id}, {e}")
        return jsonify([])
    finally:
        conn.close()


# =============================================================================
# Proposal Wizard — Section Management
# =============================================================================

@proposal_bp.route("/proposals/<prop_id>/sections/<step>/generate", methods=["POST"])
@require_role("premium")
def api_proposal_generate_section(prop_id, step):
    """Generate a proposal section using the agent with step-specific prompt and tools.

    Accepts optional body: {"instructions": "user prompt", "manual_draft": "user's own draft"}
    """
    if step not in PROPOSAL_SECTIONS:
        return jsonify({"error": f"Invalid section. Must be one of: {', '.join(PROPOSAL_SECTIONS)}"}), 400

    uid = current_uid()
    role = current_role()
    data = request.get_json(silent=True) or {}
    instructions = (data.get("instructions") or "").strip()
    manual_draft = (data.get("manual_draft") or "").strip()

    row, conn = _get_proposal_for_edit(prop_id, uid, role)
    if not row:
        conn.close() if conn else None
        return jsonify({"error": "Proposal not found or not editable"}), 404

    try:
        _update_step_status(conn, prop_id, step, "reviewing", uid, role)

        from agent.proposal_agent import generate_section
        result = generate_section(
            prop_id=prop_id,
            step=step,
            proposal_row=dict(row),
            uid=uid,
            instructions=instructions,
            manual_draft=manual_draft,
        )

        if "error" in result:
            _update_step_status(conn, prop_id, step, "draft", uid, role)
            return jsonify(result), 500

        db_field = SECTION_DB_FIELDS.get(step)
        if db_field and result.get("content"):
            content = result["content"]
            if db_field in ("toc", "logframe", "cover_page", "budget", "mne_framework", "risk_matrix"):
                if isinstance(content, str):
                    try:
                        content = json.loads(content)
                    except Exception:
                        pass
                stored = json.dumps(content) if isinstance(content, (list, dict)) else content
            else:
                stored = content if isinstance(content, str) else str(content)

            if role == "admin":
                conn.execute(f"UPDATE proposals SET {db_field} = ? WHERE id = ?", (stored, prop_id))
            else:
                conn.execute(f"UPDATE proposals SET {db_field} = ? WHERE id = ? AND uid = ?", (stored, prop_id, uid))
            conn.commit()

        _update_step_status(conn, prop_id, step, "draft", uid, role)
        _log_event(uid, "proposal_section_generated", {"prop_id": prop_id, "step": step})

        return jsonify(result)
    except Exception as e:
        logger.error(f"api_proposal_generate_section error: {prop_id}, {step}, {e}")
        _update_step_status(conn, prop_id, step, "empty", uid, role)
        return jsonify({"error": f"Section generation failed: {str(e)}"}), 500
    finally:
        conn.close()


@proposal_bp.route("/proposals/<prop_id>/upload-reference", methods=["POST"])
@require_role("premium")
def api_proposal_upload_reference(prop_id):
    """Upload reference document(s) (PDF/DOCX/TXT) for the proposal.

    Supports multiple files — text from all files is concatenated.
    Extracts text and stores as reference_text for use during section generation.
    """
    uid = current_uid()
    role = current_role()
    conn = _chats_db()
    try:
        if role == "admin":
            row = conn.execute("SELECT id FROM proposals WHERE id = ?", (prop_id,)).fetchone()
        else:
            row = conn.execute("SELECT id FROM proposals WHERE id = ? AND uid = ?", (prop_id, uid)).fetchone()
        if not row:
            return jsonify({"error": "Proposal not found"}), 404

        files = request.files.getlist("file")
        if not files or not files[0].filename:
            return jsonify({"error": "No file uploaded"}), 400

        from werkzeug.utils import secure_filename
        import tempfile
        import os as _os

        all_texts = []
        all_filenames = []
        errors = []

        for file in files:
            filename = secure_filename(file.filename or "")
            if not filename:
                continue

            ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
            if ext not in ("pdf", "docx", "doc", "txt", "md"):
                errors.append(f"{filename}: unsupported type (.{ext})")
                continue

            tmp = tempfile.NamedTemporaryFile(delete=False, suffix=f".{ext}")
            file.save(tmp.name)
            tmp.close()

            try:
                if ext == "pdf":
                    with open(tmp.name, "rb") as _f:
                        if not _f.read(5).startswith(b"%PDF"):
                            errors.append(f"{filename}: invalid PDF")
                            continue
                    from reliefweb_api.db_manager import extract_pdf_text
                    text, _pages = extract_pdf_text(tmp.name)
                elif ext in ("docx", "doc"):
                    try:
                        import docx
                        doc = docx.Document(tmp.name)
                        text = "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())
                    except ImportError:
                        errors.append(f"{filename}: DOCX parsing not available")
                        continue
                else:
                    with open(tmp.name, "r", encoding="utf-8", errors="ignore") as f:
                        text = f.read()

                text = text.strip()
                if text:
                    all_texts.append(f"--- {filename} ---\n{text}")
                    all_filenames.append(filename)
                else:
                    errors.append(f"{filename}: no text extracted")
            finally:
                _os.unlink(tmp.name)

        if not all_texts:
            return jsonify({"error": "No text could be extracted. " + "; ".join(errors)}), 400

        combined_text = "\n\n".join(all_texts)
        max_chars = 50000
        if len(combined_text) > max_chars:
            combined_text = combined_text[:max_chars] + "\n\n[... documents truncated ...]"

        combined_filename = ", ".join(all_filenames)

        # Merge with existing reference text if present
        existing_row = conn.execute("SELECT reference_text FROM proposals WHERE id = ?", (prop_id,)).fetchone()
        existing_text = ""
        if existing_row and existing_row["reference_text"]:
            existing_text = existing_row["reference_text"]
            combined_text = existing_text + "\n\n" + combined_text
            if len(combined_text) > max_chars:
                combined_text = combined_text[:max_chars] + "\n\n[... documents truncated ...]"

        if role == "admin":
            conn.execute(
                "UPDATE proposals SET reference_text = ?, reference_filename = ? WHERE id = ?",
                (combined_text, combined_filename, prop_id)
            )
        else:
            conn.execute(
                "UPDATE proposals SET reference_text = ?, reference_filename = ? WHERE id = ? AND uid = ?",
                (combined_text, combined_filename, prop_id, uid)
            )
        conn.commit()
        _log_event(uid, "proposal_reference_uploaded", {"prop_id": prop_id, "files": all_filenames})

        return jsonify({
            "message": f"{len(all_filenames)} file(s) uploaded",
            "filename": combined_filename,
            "files": all_filenames,
            "chars": len(combined_text),
            "errors": errors,
        })
    except Exception as e:
        logger.error(f"api_proposal_upload_reference error: {prop_id}, {e}")
        return jsonify({"error": "Internal server error"}), 500
    finally:
        conn.close()


@proposal_bp.route("/proposals/<prop_id>/reference", methods=["DELETE"])
@require_role("premium")
def api_proposal_delete_reference(prop_id):
    """Remove the reference document from a proposal."""
    uid = current_uid()
    role = current_role()
    conn = _chats_db()
    try:
        if role == "admin":
            row = conn.execute("SELECT id FROM proposals WHERE id = ?", (prop_id,)).fetchone()
        else:
            row = conn.execute("SELECT id FROM proposals WHERE id = ? AND uid = ?", (prop_id, uid)).fetchone()
        if not row:
            return jsonify({"error": "Proposal not found"}), 404

        if role == "admin":
            conn.execute("UPDATE proposals SET reference_text = '', reference_filename = '' WHERE id = ?", (prop_id,))
        else:
            conn.execute("UPDATE proposals SET reference_text = '', reference_filename = '' WHERE id = ? AND uid = ?", (prop_id, uid))
        conn.commit()
        return jsonify({"message": "Reference document removed"})
    except Exception as e:
        logger.error(f"api_proposal_delete_reference error: {prop_id}, {e}")
        return jsonify({"error": "Internal server error"}), 500
    finally:
        conn.close()


@proposal_bp.route("/proposals/<prop_id>/sections/<step>/revise", methods=["POST"])
@require_role("premium")
def api_proposal_revise_section(prop_id, step):
    """Revise a section via SSE streaming based on user feedback."""
    if step not in PROPOSAL_SECTIONS:
        return jsonify({"error": "Invalid section"}), 400

    uid = current_uid()
    role = current_role()
    data = request.get_json(silent=True) or {}
    feedback = (data.get("feedback") or "").strip()
    if not feedback:
        return jsonify({"error": "Feedback message is required"}), 400

    row, conn = _get_proposal_for_edit(prop_id, uid, role)
    if not row:
        conn.close() if conn else None
        return jsonify({"error": "Proposal not found or not editable"}), 404

    current_content = row[SECTION_DB_FIELDS.get(step, "")] or ""
    conn.close()

    def generate():
        try:
            from agent.proposal_agent import revise_section_stream

            yield f"data: {json.dumps({'type': 'start'})}\n\n"

            for chunk_type, chunk_data in revise_section_stream(
                prop_id=prop_id,
                step=step,
                proposal_row=dict(row),
                feedback=feedback,
                uid=uid,
            ):
                yield f"data: {json.dumps({'type': chunk_type, **chunk_data})}\n\n"

            yield f"data: {json.dumps({'type': 'done'})}\n\n"
        except Exception as e:
            logger.error(f"api_proposal_revise_section error: {prop_id}, {step}, {e}")
            yield f"data: {json.dumps({'type': 'error', 'text': 'Revision failed'})}\n\n"

    return Response(
        generate(),
        mimetype="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "X-Accel-Buffering": "no",
            "Connection": "keep-alive",
        },
    )


@proposal_bp.route("/proposals/<prop_id>/sections/<step>/approve", methods=["POST"])
@require_role("premium")
def api_proposal_approve_section(prop_id, step):
    """Approve a section, lock it, and advance to next step."""
    if step not in PROPOSAL_SECTIONS:
        return jsonify({"error": "Invalid section"}), 400

    uid = current_uid()
    role = current_role()
    conn = _chats_db()
    try:
        if role == "admin":
            row = conn.execute("SELECT current_step, step_status FROM proposals WHERE id = ?", (prop_id,)).fetchone()
        else:
            row = conn.execute(
                "SELECT current_step, step_status FROM proposals WHERE id = ? AND uid = ?",
                (prop_id, uid)
            ).fetchone()

        if not row:
            return jsonify({"error": "Proposal not found"}), 404

        step_status = json.loads(row["step_status"]) if row["step_status"] else {}
        step_status[step] = "complete"

        current_idx = PROPOSAL_SECTIONS.index(step) if step in PROPOSAL_SECTIONS else 0
        next_step = PROPOSAL_SECTIONS[current_idx + 1] if current_idx + 1 < len(PROPOSAL_SECTIONS) else "final_review"

        if next_step not in step_status or step_status[next_step] != "complete":
            step_status[next_step] = step_status.get(next_step, "empty") if next_step != step else "complete"

        is_final = step == "final_review" or current_idx == len(PROPOSAL_SECTIONS) - 1
        completed_at = _time.time() if is_final else None

        if role == "admin":
            conn.execute(
                "UPDATE proposals SET current_step = ?, step_status = ?, completed_at = COALESCE(?, completed_at) WHERE id = ?",
                (next_step, json.dumps(step_status), completed_at, prop_id)
            )
        else:
            conn.execute(
                "UPDATE proposals SET current_step = ?, step_status = ?, completed_at = COALESCE(?, completed_at) WHERE id = ? AND uid = ?",
                (next_step, json.dumps(step_status), completed_at, prop_id, uid)
            )
        conn.commit()

        _log_event(uid, "proposal_section_approved", {"prop_id": prop_id, "step": step, "next_step": next_step})

        # Run cross-section validation after every 3rd step or on final
        validation_result = None
        if (current_idx + 1) % 3 == 0 or is_final:
            try:
                from agent.validation import validate_cross_sections
                if role == "admin":
                    full_row = conn.execute("SELECT * FROM proposals WHERE id = ?", (prop_id,)).fetchone()
                else:
                    full_row = conn.execute("SELECT * FROM proposals WHERE id = ? AND uid = ?", (prop_id, uid)).fetchone()
                if full_row:
                    validation_result = validate_cross_sections(dict(full_row))
            except Exception as val_err:
                logger.warning(f"Cross-section validation failed: {val_err}")

        return jsonify({
            "message": f"Section '{step}' approved",
            "next_step": next_step,
            "step_status": step_status,
            "completed": is_final,
            "validation": validation_result,
        })
    except Exception as e:
        logger.error(f"api_proposal_approve_section error: {prop_id}, {step}, {e}")
        return jsonify({"error": "Internal server error"}), 500
    finally:
        conn.close()


@proposal_bp.route("/proposals/<prop_id>/sections/<step>", methods=["PUT"])
@require_role("premium")
def api_proposal_update_section(prop_id, step):
    """Manually update a section's content (user edits)."""
    if step not in PROPOSAL_SECTIONS:
        return jsonify({"error": "Invalid section"}), 400

    uid = current_uid()
    role = current_role()
    data = request.get_json(silent=True) or {}
    content = data.get("content")

    if content is None:
        return jsonify({"error": "content field is required"}), 400

    db_field = SECTION_DB_FIELDS.get(step)
    if not db_field:
        return jsonify({"error": "No database field for this section"}), 400

    conn = _chats_db()
    try:
        if role == "admin":
            row = conn.execute("SELECT id FROM proposals WHERE id = ?", (prop_id,)).fetchone()
        else:
            row = conn.execute("SELECT id FROM proposals WHERE id = ? AND uid = ?", (prop_id, uid)).fetchone()

        if not row:
            return jsonify({"error": "Proposal not found"}), 404

        if db_field in ("toc", "logframe", "cover_page", "budget", "mne_framework", "risk_matrix"):
            if isinstance(content, str):
                try:
                    parsed = json.loads(content)
                    stored = json.dumps(parsed)
                except Exception:
                    stored = content
            else:
                stored = json.dumps(content)
        else:
            stored = content if isinstance(content, str) else str(content)

        if role == "admin":
            conn.execute(f"UPDATE proposals SET {db_field} = ? WHERE id = ?", (stored, prop_id))
        else:
            conn.execute(f"UPDATE proposals SET {db_field} = ? WHERE id = ? AND uid = ?", (stored, prop_id, uid))
        conn.commit()

        return jsonify({"message": "Section updated"})
    except Exception as e:
        logger.error(f"api_proposal_update_section error: {prop_id}, {step}, {e}")
        return jsonify({"error": "Internal server error"}), 500
    finally:
        conn.close()


# =============================================================================
# Donor templates
# =============================================================================

@proposal_bp.route("/donor-templates", methods=["GET"])
@require_auth
def api_donor_templates():
    """List available donor templates."""
    try:
        from agent.donor_templates import list_templates
        return jsonify(list_templates())
    except Exception as e:
        logger.error(f"api_donor_templates error: {e}")
        return jsonify([]), 500




# =============================================================================
# Proposal export
# =============================================================================

@proposal_bp.route("/proposals/<prop_id>/export", methods=["POST"])
@require_auth
def api_proposal_export(prop_id):
    """Compile all sections into a full markdown proposal."""
    uid = current_uid()
    role = current_role()
    conn = _chats_db()
    try:
        if role == "admin":
            row = conn.execute("SELECT * FROM proposals WHERE id = ?", (prop_id,)).fetchone()
        else:
            row = conn.execute("SELECT * FROM proposals WHERE id = ? AND uid = ?", (prop_id, uid)).fetchone()

        if not row:
            allowed = conn.execute(
                "SELECT completed_at FROM proposals WHERE id = ?", (prop_id,)
            ).fetchone()
            if not allowed or not allowed["completed_at"]:
                return jsonify({"error": "Proposal not found"}), 404
            row = conn.execute("SELECT * FROM proposals WHERE id = ?", (prop_id,)).fetchone()
            if not row:
                return jsonify({"error": "Proposal not found"}), 404

        # Helper for parsing JSON safely
        def safe_json(val, default):
            if not val or val in ("", "{}", "[]", "null"): return default
            try: return json.loads(val)
            except Exception: return val if isinstance(val, (dict, list, str)) else default

        cover = safe_json(row.get("cover_page"), {})
        toc = safe_json(row.get("toc"), []) or safe_json(row.get("toc_nodes"), [])
        logframe = safe_json(row.get("logframe"), {}) or safe_json(row.get("logframe_data"), {})
        budget = safe_json(row.get("budget"), {}) or safe_json(row.get("budget_details"), {})
        mne = safe_json(row.get("mne_framework"), {}) or safe_json(row.get("mne_plan"), [])
        risks = safe_json(row.get("risk_matrix"), []) or safe_json(row.get("risk_details"), [])
        b_data = safe_json(row.get("beneficiary_data"), {})
        pinned_sources = safe_json(row.get("pinned_sources"), [])

        try:
            themes_list = json.loads(row["themes"])
        except Exception:
            themes_list = [t.strip() for t in row["themes"].split(",") if t.strip()]

        md_parts = []
        md_parts.append(f"# {row['title']}\n")
        
        meta_lines = [
            f"**Country of Operation:** {row['country']}",
            f"**Target Donor:** {row['donor']}",
            f"**Focus Sector / Event:** {row['event']}",
            f"**Themes:** {', '.join(themes_list) if themes_list else 'Humanitarian Response'}"
        ]
        if row.get("date_from") or row.get("date_to"):
            meta_lines.append(f"**Timeline:** {row.get('date_from', '')} to {row.get('date_to', '')}")
        md_parts.append("  \n".join(meta_lines) + "\n")

        if cover and isinstance(cover, dict) and any(cover.values()):
            md_parts.append("\n## Project Overview\n")
            for k, v in cover.items():
                if v:
                    label = k.replace('_', ' ').title()
                    md_parts.append(f"**{label}:** {v}  \n")
        elif isinstance(cover, str):
            md_parts.append(f"\n## Project Overview\n\n{cover}\n")

        if b_data and isinstance(b_data, dict) and any(b_data.values()):
            md_parts.append("\n## Target Beneficiaries\n")
            md_parts.append("| Category / Group | Direct Beneficiaries | Indirect Beneficiaries | Details |\n|---|---|---|---|\n")
            direct = b_data.get("direct", {})
            indirect = b_data.get("indirect", "N/A")
            if isinstance(direct, dict):
                md_parts.append(f"| Women | {direct.get('women', '-')} | - | Targeted female beneficiaries |\n")
                md_parts.append(f"| Men | {direct.get('men', '-')} | - | Targeted male beneficiaries |\n")
                md_parts.append(f"| Children | {direct.get('children', '-')} | - | Targeted boys & girls |\n")
                total_dir = b_data.get("total_direct") or direct.get("total") or "Sum of categories"
                md_parts.append(f"| **TOTAL DIRECT** | **{total_dir}** | **{indirect}** | Total projected reach |\n")
            elif isinstance(b_data, list):
                for item in b_data:
                    if isinstance(item, dict):
                        md_parts.append(f"| {item.get('group', 'Group')} | {item.get('direct', '-')} | {item.get('indirect', '-')} | {item.get('notes', '')} |\n")
        elif isinstance(b_data, str):
            md_parts.append(f"\n## Target Beneficiaries\n\n{b_data}\n")

        if row.get("background"):
            md_parts.append(f"\n## Context & Background\n\n{row['background']}\n")
        if row.get("needs_assessment"):
            md_parts.append(f"\n## Needs Assessment\n\n{row['needs_assessment']}\n")

        if toc:
            md_parts.append("\n## Theory of Change\n")
            if isinstance(toc, list):
                md_parts.append("| Level | Summary / Intervention Node | Causal Relationship / Notes |\n|---|---|---|\n")
                for node in toc:
                    if isinstance(node, dict):
                        lvl = node.get("level", "Node").title()
                        txt = node.get("text") or node.get("label") or str(node)
                        parents = ", ".join(node.get("parent_ids", []))
                        note = f"Contributes to: {parents}" if parents else "Core Intervention Step"
                        md_parts.append(f"| {lvl} | {txt} | {note} |\n")
                    else:
                        md_parts.append(f"| Step | {str(node)} | Strategic Flow |\n")
            elif isinstance(toc, str):
                md_parts.append(f"\n{toc}\n")

        if logframe:
            md_parts.append("\n## Logical Framework (Logframe)\n")
            if isinstance(logframe, dict):
                md_parts.append("| Level | Objective / Summary | Performance Indicators | Means of Verification | Risks & Assumptions |\n|---|---|---|---|---|\n")
                levels = [("goal", "Overall Goal / Impact"), ("outcomes", "Outcomes"), ("outputs", "Outputs"), ("activities", "Activities")]
                has_levels = False
                for key, label in levels:
                    val = logframe.get(key)
                    if val:
                        has_levels = True
                        if isinstance(val, list):
                            for idx, item in enumerate(val):
                                if isinstance(item, dict):
                                    md_parts.append(f"| {label} #{idx+1} | {item.get('text', item.get('description', ''))} | {item.get('indicators', '-')} | {item.get('verification', '-')} | {item.get('assumptions', '-')} |\n")
                                else:
                                    md_parts.append(f"| {label} #{idx+1} | {str(item)} | Key indicator TBD | Progress reports | Assumptions valid |\n")
                        elif isinstance(val, dict):
                            md_parts.append(f"| {label} | {val.get('text', val.get('description', ''))} | {val.get('indicators', '-')} | {val.get('verification', '-')} | {val.get('assumptions', '-')} |\n")
                        else:
                            md_parts.append(f"| {label} | {str(val)} | Standard Indicators | Verification Records | Core Assumptions |\n")
                if not has_levels:
                    for k, v in logframe.items():
                        lbl = k.replace('_', ' ').title()
                        md_parts.append(f"| {lbl} | {str(v)} | Specified in M&E | Field Audits | Low Risk |\n")
            elif isinstance(logframe, str):
                md_parts.append(f"\n{logframe}\n")

        if row.get("methodology"):
            md_parts.append(f"\n## Implementation Methodology\n\n{row['methodology']}\n")

        if budget:
            md_parts.append("\n## Budget Summary\n")
            if isinstance(budget, (dict, list)):
                md_parts.append("| Category | Line Item / Description | Amount ($) | Share / Notes |\n|---|---|---|---|\n")
                lines = budget.get("lines", []) if isinstance(budget, dict) else budget
                total_val = budget.get("total") if isinstance(budget, dict) else None

                calc_total = 0.0
                for line in lines:
                    if isinstance(line, dict):
                        cat = line.get("category", "General").title()
                        desc = line.get("description", line.get("item", "-"))
                        amt = line.get("amount", 0)
                        try:
                            amt_num = float(str(amt).replace('$', '').replace(',', ''))
                            calc_total += amt_num
                            amt_str = f"${amt_num:,.2f}"
                        except Exception:
                            amt_str = str(amt)
                        md_parts.append(f"| {cat} | {desc} | {amt_str} | Standard operational expense |\n")

                disp_total = total_val if total_val else (f"${calc_total:,.2f}" if calc_total > 0 else "N/A")
                md_parts.append(f"| **TOTAL PROJECT BUDGET** | **Grand Total** | **{disp_total}** | **100% Allocated** |\n")
            elif isinstance(budget, str):
                md_parts.append(f"\n{budget}\n")

        if mne:
            md_parts.append("\n## Monitoring & Evaluation Framework\n")
            if isinstance(mne, (dict, list)):
                md_parts.append("| Indicator Name | Baseline | Target | Means of Verification | Frequency & Lead |\n|---|---|---|---|---|\n")
                indicators = mne.get("indicators", []) if isinstance(mne, dict) else mne
                for ind in indicators:
                    if isinstance(ind, dict):
                        name = ind.get("name", ind.get("indicator", "Indicator"))
                        base = ind.get("baseline", "0")
                        tgt = ind.get("target", "100%")
                        src = ind.get("source", ind.get("verification", "Field Reports"))
                        freq = ind.get("frequency", ind.get("owner", "Monthly / M&E Officer"))
                        md_parts.append(f"| {name} | {base} | {tgt} | {src} | {freq} |\n")
            elif isinstance(mne, str):
                md_parts.append(f"\n{mne}\n")

        if risks:
            md_parts.append("\n## Risk Matrix & Mitigation Strategy\n")
            if isinstance(risks, list):
                md_parts.append("| Risk Category / Threat | Severity / Category | Probability | Impact | Mitigation Strategy |\n|---|---|---|---|---|\n")
                for r in risks:
                    if isinstance(r, dict):
                        risk_txt = r.get("risk", r.get("name", "Identified Risk"))
                        cat = r.get("category", "Operational").title()
                        prob = r.get("probability", "Medium").title()
                        imp = r.get("impact", "Medium").title()
                        mit = r.get("mitigation", r.get("strategy", "Continuous monitoring"))
                        md_parts.append(f"| {risk_txt} | {cat} | {prob} | {imp} | {mit} |\n")
            elif isinstance(risks, str):
                md_parts.append(f"\n{risks}\n")

        if pinned_sources:
            md_parts.append("\n## Data Sources & Operational Evidence\n")
            if isinstance(pinned_sources, list):
                md_parts.append("| Source Title | Reference URL / ID | Context Snippet |\n|---|---|---|\n")
                for src in pinned_sources:
                    if isinstance(src, dict):
                        stitle = src.get("title", "ReliefWeb/HDX Report")
                        surl = src.get("url", "#")
                        ssnip = src.get("snippet", "-")[:120] + "..." if len(src.get("snippet", "")) > 120 else src.get("snippet", "-")
                        md_parts.append(f"| {stitle} | [{surl}]({surl}) | {ssnip} |\n")
            elif isinstance(pinned_sources, str):
                md_parts.append(f"\n{pinned_sources}\n")


        full_md = "\n".join(md_parts)
        _log_event(uid, "proposal_exported", {"prop_id": prop_id})

        return jsonify({
            "markdown": full_md,
            "title": row["title"],
            "filename": f"proposal_{prop_id}.md",
        })
    except Exception as e:
        logger.error(f"api_proposal_export error: {prop_id}, {e}")
        return jsonify({"error": "Internal server error"}), 500
    finally:
        conn.close()



# ═══════════════════════════════════════════════════════════════════════════
# PROPOSAL REVIEW — Structured AI feedback (replaces advisor chat)
# ═══════════════════════════════════════════════════════════════════════════

@proposal_bp.route("/proposals/<prop_id>/review", methods=["POST"])
@require_role("premium")
def api_proposal_review(prop_id):
    """AI-powered structured review of entire proposal.
    
    Returns section-by-section analysis with scores, priorities, and suggestions.
    Uses gemini-2.5-flash for fast, high-quality review.
    """
    uid = current_uid()
    conn = _chats_db()
    try:
        row = conn.execute(
            "SELECT * FROM proposals WHERE id = ? AND uid = ?",
            (prop_id, uid)
        ).fetchone()

        if not row:
            return jsonify({"error": "Proposal not found"}), 404

        # Collect all section content for review
        sections_content = {}
        field_map = {
            "cover": "cover_page", "background": "background",
            "needs_assessment": "needs_assessment", "toc": "toc",
            "logframe": "logframe", "methodology": "methodology",
            "budget": "budget", "mne_framework": "mne_framework",
            "risk_matrix": "risk_matrix", "sustainability": "sustainability",
            "coordination": "coordination", "final_review": "narrative",
        }

        for step, field in field_map.items():
            content = row[field]
            if content and content not in ("", "{}", "[]", "null"):
                try:
                    parsed = json.loads(content)
                    if isinstance(parsed, dict) and "content" in parsed:
                        sections_content[step] = {
                            "content": parsed["content"][:2000],
                            "sources": parsed.get("sources", [])
                        }
                    else:
                        sections_content[step] = {"content": json.dumps(parsed, indent=2)[:2000]}
                except (json.JSONDecodeError, TypeError):
                    sections_content[step] = {"content": str(content)[:2000]}

        # Check step_status for skipped sections
        try:
            step_status = json.loads(row["step_status"]) if row["step_status"] else {}
        except (json.JSONDecodeError, TypeError):
            step_status = {}

        # Build the review prompt with all section content
        from agent.proposal_prompts import REVIEW_SYSTEM_PROMPT
        section_texts = []
        for step, data in sections_content.items():
            status = step_status.get(step, "pending")
            section_texts.append(f"## {step.replace('_', ' ').title()} [status: {status}]\n{data['content']}")

        # Add skipped/empty sections
        for step, field in field_map.items():
            if step not in sections_content:
                status = step_status.get(step, "pending")
                section_texts.append(f"## {step.replace('_', ' ').title()} [status: {status}]\n(empty)")

        full_proposal_text = "\n\n---\n\n".join(section_texts)

        metadata = f"Country: {row['country']}\nEvent/Crisis: {row['event']}\nDonor: {row['donor']}"

        user_message = f"{metadata}\n\n--- PROPOSAL CONTENT ---\n\n{full_proposal_text}"

        # Use Gemini 2.5 Flash for review (fast, high quality, good JSON)
        from langchain_openai import ChatOpenAI
        from config import config as _cfg

        review_model = os.environ.get("REVIEW_MODEL", "google/gemini-2.5-flash")
        model = ChatOpenAI(
            model=review_model,
            base_url=_cfg._LLM_BASE_URL,
            api_key=_cfg._LLM_API_KEY,
            temperature=0.3,
            max_tokens=4096,
            timeout=90,
        )

        from langchain_core.messages import HumanMessage, SystemMessage
        messages = [
            SystemMessage(content=REVIEW_SYSTEM_PROMPT),
            HumanMessage(content=user_message),
        ]

        response = model.invoke(messages)
        review_text = response.content.strip()

        # Parse JSON response
        if review_text.startswith("```json"):
            review_text = review_text[7:]
        if review_text.startswith("```"):
            review_text = review_text[3:]
        if review_text.endswith("```"):
            review_text = review_text[:-3]

        try:
            review_data = json.loads(review_text)
        except json.JSONDecodeError:
            review_data = {
                "sections": [],
                "overall_score": 0,
                "overall_feedback": review_text[:500],
                "high_priority": [],
                "medium_priority": [],
                "strengths": [],
                "suggested_actions": [],
                "raw_response": review_text
            }

        # Collect sources from sections
        all_sources = []
        for step, data in sections_content.items():
            if isinstance(data, dict):
                for src in data.get("sources", []):
                    if src not in all_sources:
                        all_sources.append(src)
        review_data["sources"] = all_sources

        # Save review to DB
        conn.execute(
            "UPDATE proposals SET review = ? WHERE id = ? AND uid = ?",
            (json.dumps(review_data), prop_id, uid)
        )
        conn.commit()

        _log_event(uid, "proposal_reviewed", {"prop_id": prop_id, "score": review_data.get("overall_score", 0)})

        return jsonify(review_data)

    except Exception as e:
        logger.error(f"api_proposal_review error: {prop_id}, {e}")
        return jsonify({"error": f"Review failed: {str(e)}"}), 500
    finally:
        conn.close()